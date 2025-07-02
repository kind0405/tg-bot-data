import os
import logging
import requests
import io

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload
from telegram import Update
from telegram.ext import Application, MessageHandler, filters, ContextTypes
from PyPDF2 import PdfReader
from docx import Document
import openpyxl

# === ТВОИ ДАННЫЕ ===
TELEGRAM_TOKEN = '7746119786:AAGm0uWy-urxACu9Q9w0lP9HQ6v610K6Vcg'
CLOUDFLARE_API_TOKEN = 'uvxkh4dYn8n5ntavSNPdNKG8qRHHujOlxPAV6zVz'
CLOUDFLARE_AI_URL = 'https://api.cloudflare.com/client/v4/accounts/28d3650e64ec7f2490b9a6dd14e6b659/ai/run/@cf/meta/llama-2-7b-chat-int8'
FOLDER_ID = '1SDBfV-2Zk7lriKUsgRSS6wWnyC2O7ZX0'
SERVICE_ACCOUNT_FILE = 'credentials.json'

# === ЛОГГИРОВАНИЕ ===
logging.basicConfig(level=logging.INFO)

# === GOOGLE DRIVE ===
SCOPES = ['https://www.googleapis.com/auth/drive.readonly']
creds = service_account.Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=SCOPES)
drive_service = build('drive', 'v3', credentials=creds)

# === ЧТЕНИЕ ДОКУМЕНТОВ ===
def extract_text_from_file(file_id, mime_type, name):
    try:
        if mime_type.startswith('application/vnd.google-apps.document'):
            request = drive_service.files().export_media(fileId=file_id, mimeType='text/plain')
        else:
            request = drive_service.files().get_media(fileId=file_id)

        fh = io.BytesIO()
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        fh.seek(0)

        if mime_type == 'application/pdf':
            reader = PdfReader(fh)
            return "\n".join(page.extract_text() or '' for page in reader.pages)

        elif mime_type.endswith('document') or mime_type == 'text/plain':
            return fh.read().decode('utf-8', errors='ignore')

        elif mime_type.endswith('wordprocessingml.document'):
            doc = Document(fh)
            return '\n'.join(p.text for p in doc.paragraphs)

        elif mime_type.endswith('spreadsheetml.sheet'):
            wb = openpyxl.load_workbook(fh, data_only=True)
            text = ''
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += ' | '.join(str(cell) if cell else '' for cell in row) + '\n'
            return text

        else:
            logging.warning(f"Пропущен неподдерживаемый файл: {name}")
            return ""

    except HttpError as e:
        logging.warning(f"Не удалось прочитать файл {name}: {e}")
        return ""

def load_all_documents():
    query = f"'{FOLDER_ID}' in parents and trashed = false"
    results = drive_service.files().list(q=query, fields="files(id, name, mimeType)").execute()
    files = results.get('files', [])

    all_texts = []
    for file in files:
        text = extract_text_from_file(file['id'], file['mimeType'], file['name'])
        if text:
            all_texts.append(f"==== {file['name']} ====\n{text}")
    return "\n\n".join(all_texts)

documents_text = load_all_documents()

# === CLOUDFLARE AI ===
def ask_ai(prompt):
    headers = {
        "Authorization": f"Bearer {CLOUDFLARE_API_TOKEN}",
        "Content-Type": "application/json"
    }
    data = {
        "messages": [
            {"role": "system", "content": "Ты помощник, который отвечает на вопросы по внутренним документам компании. Отвечай понятно и по существу."},
            {"role": "user", "content": prompt}
        ]
    }
    try:
        response = requests.post(CLOUDFLARE_AI_URL, headers=headers, json=data)
        response.raise_for_status()
        result = response.json()
        return result.get("result", {}).get("response", "AI не ответил.")
    except Exception as e:
        logging.error(f"Ошибка при запросе к Cloudflare AI: {e}")
        return "Ошибка при обращении к AI."

# === ОБРАБОТКА СООБЩЕНИЙ ОТ ПОЛЬЗОВАТЕЛЯ ===
async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_input = update.message.text
    prompt = f"Вопрос: {user_input}\n\nВот документы компании:\n{documents_text}"
    response = ask_ai(prompt)
    await update.message.reply_text(response[:4000])  # ограничение Telegram

# === СТАРТ БОТА ===
def main():
    app = Application.builder().token(TELEGRAM_TOKEN).build()
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    app.run_polling()

if __name__ == '__main__':
    main()

